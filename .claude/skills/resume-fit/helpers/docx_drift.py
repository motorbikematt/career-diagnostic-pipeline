"""Detect drift between a resume's markdown source and its rendered .docx.

`render_docx.py` is one-way: markdown in, .docx out. When the user edits the
.docx directly in Word (a normal thing to do), the markdown silently goes stale,
and the next pipeline pass works from it and loses those edits. This happened
repeatedly in live runs.

This helper compares the VISIBLE TEXT of both files, paragraph by paragraph,
after the same markdown stripping render_docx.py applies. Formatting-only
changes (bold, font) are not drift; any wording change is. Whitespace is
collapsed and non-breaking spaces are normalized, but characters Word
autocorrect inserts (curly quotes, en dashes) ARE reported: they are real
changes, and ats_chars.py treats them as ATS hazards.

`--pull <out.md>` rebuilds markdown from the .docx (headings, bullets, bold,
italic, mirroring render_docx.py's styles) so the user's Word edits can become
the new source of truth. Review the pulled file before replacing the original.
"""
from __future__ import annotations

import difflib
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

_INLINE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*")
_WS = re.compile(r"\s+")


def _norm(text: str) -> str:
    return _WS.sub(" ", text.replace(" ", " ")).strip()


def md_visible_lines(md_text: str) -> list:
    """Visible paragraph texts, as render_docx.py would lay them out."""
    out = []
    for raw in md_text.splitlines():
        line = raw.rstrip()
        if not line.strip() or line.strip() == "---":
            continue
        for prefix in ("# ", "## ", "### "):
            if line.startswith(prefix):
                line = line[len(prefix):]
                break
        else:
            if re.match(r"^[-*] ", line):
                line = line[2:]
        line = _INLINE.sub(lambda m: m.group(1) or m.group(2), line)
        out.append(_norm(line))
    return out


def docx_visible_lines(docx_path) -> list:
    doc = Document(str(docx_path))
    return [_norm(p.text) for p in doc.paragraphs if _norm(p.text)]


def compare(md_path, docx_path) -> dict:
    md_path, docx_path = Path(md_path), Path(docx_path)
    md_lines = md_visible_lines(md_path.read_text(encoding="utf-8"))
    docx_lines = docx_visible_lines(docx_path)
    diff = list(difflib.unified_diff(
        md_lines, docx_lines, fromfile=md_path.name, tofile=docx_path.name,
        lineterm="", n=0,
    ))
    return {
        "in_sync": md_lines == docx_lines,
        "docx_newer": docx_path.stat().st_mtime > md_path.stat().st_mtime,
        "only_in_md": [l[1:] for l in diff if l.startswith("-") and not l.startswith("---")],
        "only_in_docx": [l[1:] for l in diff if l.startswith("+") and not l.startswith("+++")],
        "diff": diff,
    }


def _runs_md(paragraph) -> str:
    """Inline markdown for a paragraph, merging adjacent same-format runs
    (Word splits runs arbitrarily, e.g. after spell-check)."""
    groups: list = []
    for r in paragraph.runs:
        if not r.text:
            continue
        fmt = (bool(r.bold), bool(r.italic))
        if groups and groups[-1][0] == fmt:
            groups[-1][1].append(r.text)
        else:
            groups.append((fmt, [r.text]))
    parts = []
    for (bold, italic), texts in groups:
        t = "".join(texts)
        core = t.strip()
        if not core or not (bold or italic):
            parts.append(t)
            continue
        mark = "**" if bold else "*"
        lead = t[: len(t) - len(t.lstrip())]
        trail = t[len(t.rstrip()):]
        parts.append(f"{lead}{mark}{core}{mark}{trail}")
    return "".join(parts).replace(" ", " ").strip()


def _heading_level(paragraph):
    runs = [r for r in paragraph.runs if r.text.strip()]
    if not runs or not all(r.bold for r in runs):
        return None
    size = runs[0].font.size
    if size == Pt(20):
        return 1
    if size == Pt(13):
        return 2
    if paragraph.paragraph_format.space_before == Pt(6):
        return 3
    return None


def docx_to_md(docx_path) -> str:
    doc = Document(str(docx_path))
    out: list = []
    for p in doc.paragraphs:
        text = _norm(p.text)
        if not text:
            continue
        level = _heading_level(p)
        if level:
            if out:
                out.append("")
            out.append(f"{'#' * level} {text}")
            if level < 3:
                out.append("")
        elif p.style is not None and p.style.name.startswith("List"):
            out.append(f"- {_runs_md(p)}")
        else:
            out.append(_runs_md(p))
    return "\n".join(out) + "\n"


if __name__ == "__main__":
    import argparse
    import json

    from console import utf8_stdout

    utf8_stdout()
    ap = argparse.ArgumentParser(description="Detect md vs docx drift.")
    ap.add_argument("md")
    ap.add_argument("docx")
    ap.add_argument("--pull", metavar="OUT_MD",
                    help="write markdown rebuilt from the docx to this file")
    args = ap.parse_args()

    result = compare(args.md, args.docx)
    if args.pull:
        Path(args.pull).write_text(docx_to_md(args.docx), encoding="utf-8")
        result["pulled_to"] = args.pull
    print(json.dumps({k: v for k, v in result.items() if k != "diff"},
                     indent=2, ensure_ascii=False))
    if not result["in_sync"]:
        print("DRIFT: the .docx and .md differ. Reconcile before editing further:")
        print("\n".join(result["diff"]))
        raise SystemExit(1)
