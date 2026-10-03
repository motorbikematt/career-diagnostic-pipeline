"""Contact-details location check for a .docx resume (deterministic).

Greenhouse documents contact details placed only in a header, footer or text box
as a parse failure [S1]. `render_docx.py` cannot emit headers or footers, so the
risk is in source resumes the user supplies or edits in Word.

The check splits a .docx into regions (body paragraphs, table cells, headers,
footers, text boxes), finds email and phone tokens (and an optional full name),
and FAILS when a token appears only outside the body text. Detection is exact
pattern matching, so Python owns it; the helper reports, it never edits.

`docx_text_by_region()` is public so the post-approval extraction round trip can
reuse it. `docx_drift.py` is intentionally left alone.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE = re.compile(r"(?:\+?\d{1,2}[\s.-]?)?(?:\(\d{3}\)|\d{3})[\s.-]?\d{3}[\s.-]?\d{4}")

# Regions whose text counts as "in the body" for extraction purposes.
BODY_REGIONS = ("body", "table")


def _textbox_lines(element) -> list:
    """Text of every text box (w:txbxContent) under `element`, one string per paragraph."""
    out = []
    for box in element.iter(f"{_W}txbxContent"):
        for p in box.iter(f"{_W}p"):
            text = "".join(t.text or "" for t in p.iter(f"{_W}t")).strip()
            if text:
                out.append(text)
    return out


def _header_footer_parts(doc) -> tuple:
    headers, footers = [], []
    for s in doc.sections:
        headers += [s.header, s.first_page_header, s.even_page_header]
        footers += [s.footer, s.first_page_footer, s.even_page_footer]
    return headers, footers


def _part_lines(part) -> list:
    lines = [p.text for p in part.paragraphs if p.text.strip()]
    for t in part.tables:
        lines += [c.text for row in t.rows for c in row.cells if c.text.strip()]
    lines += _textbox_lines(part._element)
    return lines


def docx_text_by_region(path) -> dict:
    """Return {region: [text lines]} for body, table, header, footer and textbox."""
    doc = Document(str(path))
    headers, footers = _header_footer_parts(doc)
    return {
        "body": [p.text for p in doc.paragraphs if p.text.strip()],
        "table": [c.text for t in doc.tables for row in t.rows for c in row.cells
                  if c.text.strip()],
        "header": [l for h in headers for l in _part_lines(h)],
        "footer": [l for f in footers for l in _part_lines(f)],
        "textbox": _textbox_lines(doc.element.body),
    }


def _find(kind: str, pattern, regions: dict) -> dict:
    found_in = []
    tokens = set()
    for region, lines in regions.items():
        hits = [m.group(0) for line in lines for m in pattern.finditer(line)]
        if hits:
            found_in.append(region)
            tokens.update(hits)
    return {
        "kind": kind,
        "tokens": sorted(tokens),
        "found_in": found_in,
        "in_body": any(r in BODY_REGIONS for r in found_in),
    }


def _find_name(name: str, regions: dict) -> dict:
    needle = name.lower()
    found_in = [r for r, lines in regions.items() if any(needle in l.lower() for l in lines)]
    return {
        "kind": "name",
        "tokens": [name] if found_in else [],
        "found_in": found_in,
        "in_body": any(r in BODY_REGIONS for r in found_in),
    }


def check_regions(regions: dict, name: str | None = None) -> dict:
    findings = [_find("email", EMAIL, regions), _find("phone", PHONE, regions)]
    if name:
        findings.append(_find_name(name, regions))
    failures, warnings = [], []
    for f in findings:
        if f["found_in"] and not f["in_body"]:
            failures.append(f"{f['kind']} appears only in {', '.join(f['found_in'])}; "
                            "move it into the document body")
        elif not f["found_in"]:
            warnings.append(f"{f['kind']} not found anywhere in the document")
    return {"ok": not failures, "findings": findings,
            "failures": failures, "warnings": warnings}


def check_file(path, name: str | None = None) -> dict:
    return check_regions(docx_text_by_region(Path(path)), name=name)


if __name__ == "__main__":
    import argparse
    import json
    import sys

    from console import utf8_stdout

    utf8_stdout()
    ap = argparse.ArgumentParser(
        description="Fail if contact details appear only in a header, footer or text box.")
    ap.add_argument("docx")
    ap.add_argument("--name", help="full name to look for as well as email and phone")
    args = ap.parse_args()

    result = check_file(args.docx, name=args.name)
    print(json.dumps(result, indent=2, ensure_ascii=False))
    if not result["ok"]:
        print("CONTACT DETAILS OUTSIDE THE BODY: " + "; ".join(result["failures"]))
        sys.exit(1)
