"""Contact-location check on synthetic .docx files (TODO #16)."""
from docx import Document
from docx.oxml import parse_xml

import contact_check

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
EMAIL = "pat.example@example.com"
PHONE = "555-123-4567"


def _textbox_xml(text: str):
    return parse_xml(
        f'<w:p xmlns:w="{_W_NS}"><w:r><w:pict><w:txbxContent>'
        f'<w:p><w:r><w:t>{text}</w:t></w:r></w:p>'
        f'</w:txbxContent></w:pict></w:r></w:p>'
    )


def _save(doc, tmp_path):
    p = tmp_path / "r.docx"
    doc.save(str(p))
    return p


def test_body_contact_passes(tmp_path):
    d = Document()
    d.add_paragraph("Pat Example")
    d.add_paragraph(f"{EMAIL} | {PHONE}")
    r = contact_check.check_file(_save(d, tmp_path), name="Pat Example")
    assert r["ok"] is True
    assert r["failures"] == []


def test_header_only_contact_fails(tmp_path):
    d = Document()
    d.add_paragraph("Experience")
    d.sections[0].header.paragraphs[0].text = f"{EMAIL} | {PHONE}"
    r = contact_check.check_file(_save(d, tmp_path))
    assert r["ok"] is False
    kinds = {f["kind"]: f for f in r["findings"]}
    assert kinds["email"]["found_in"] == ["header"]
    assert kinds["email"]["in_body"] is False


def test_footer_only_contact_fails(tmp_path):
    d = Document()
    d.add_paragraph("Experience")
    d.sections[0].footer.paragraphs[0].text = EMAIL
    r = contact_check.check_file(_save(d, tmp_path))
    assert r["ok"] is False


def test_textbox_only_contact_fails(tmp_path):
    d = Document()
    d.add_paragraph("Experience")
    d.element.body.insert(0, _textbox_xml(f"{EMAIL} {PHONE}"))
    regions = contact_check.docx_text_by_region(_save(d, tmp_path))
    assert any(EMAIL in l for l in regions["textbox"])
    assert not any(EMAIL in l for l in regions["body"])
    assert contact_check.check_regions(regions)["ok"] is False


def test_contact_in_both_header_and_body_passes(tmp_path):
    d = Document()
    d.add_paragraph(f"{EMAIL} {PHONE}")
    d.sections[0].header.paragraphs[0].text = EMAIL
    assert contact_check.check_file(_save(d, tmp_path))["ok"] is True


def test_missing_contact_warns_not_fails(tmp_path):
    d = Document()
    d.add_paragraph("Experience")
    r = contact_check.check_file(_save(d, tmp_path))
    assert r["ok"] is True
    assert len(r["warnings"]) == 2


def test_name_only_in_header_fails(tmp_path):
    d = Document()
    d.add_paragraph(f"{EMAIL} {PHONE}")
    d.sections[0].header.paragraphs[0].text = "Pat Example"
    r = contact_check.check_file(_save(d, tmp_path), name="Pat Example")
    assert r["ok"] is False
    assert "name appears only in header" in r["failures"][0]
